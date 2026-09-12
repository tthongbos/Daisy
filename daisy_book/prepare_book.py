from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document

from .text import looks_like_chapter_heading, normalize_text


def safe_filename(index: int, title: str) -> str:
    slug = re.sub(r"[^0-9A-Za-zÀ-ỹ]+", "_", title, flags=re.UNICODE).strip("_")
    if not slug:
        slug = f"chapter_{index:02d}"
    return f"chapter_{index:02d}_{slug[:60]}.txt"


def prepare(source: Path, clean_dir: Path, text_dir: Path) -> dict:
    if not source.exists():
        raise FileNotFoundError(f"Source DOCX not found: {source}")

    clean_dir.mkdir(parents=True, exist_ok=True)
    text_dir.mkdir(parents=True, exist_ok=True)

    src = Document(source)
    cleaned = Document()

    # Remove the empty paragraph created by a fresh Document only when safe.
    if len(cleaned.paragraphs) == 1 and not cleaned.paragraphs[0].text:
        p = cleaned.paragraphs[0]._element
        p.getparent().remove(p)

    records: list[dict[str, str]] = []
    for paragraph in src.paragraphs:
        text = normalize_text(paragraph.text)
        style_name = paragraph.style.name if paragraph.style else "Normal"
        records.append({"text": text, "style": style_name})

        new_p = cleaned.add_paragraph(text)
        try:
            new_p.style = style_name
        except KeyError:
            pass

    clean_path = clean_dir / "book_clean.docx"
    cleaned.save(clean_path)

    chapter_starts = [
        i for i, item in enumerate(records)
        if item["text"] and looks_like_chapter_heading(item["text"], item["style"])
    ]

    manifest: dict = {
        "source": str(source),
        "clean_docx": str(clean_path),
        "chapter_detection": "heading-or-regex",
        "chapters": [],
    }

    if not chapter_starts:
        full_text = "\n\n".join(item["text"] for item in records if item["text"])
        target = text_dir / "full_book.txt"
        target.write_text(full_text + "\n", encoding="utf-8")
        manifest["chapter_detection"] = "none"
        manifest["chapters"].append({"index": 1, "title": "Full book", "file": target.name})
        print("Warning: no chapter headings detected. Wrote build/text/full_book.txt")
    else:
        for chapter_index, start in enumerate(chapter_starts, start=1):
            end = chapter_starts[chapter_index] if chapter_index < len(chapter_starts) else len(records)
            title = records[start]["text"]
            body = [item["text"] for item in records[start:end] if item["text"]]
            filename = safe_filename(chapter_index, title)
            target = text_dir / filename
            target.write_text("\n\n".join(body) + "\n", encoding="utf-8")
            manifest["chapters"].append({"index": chapter_index, "title": title, "file": filename})

    manifest_path = text_dir / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return manifest


def main() -> int:
    parser = argparse.ArgumentParser(description="Normalize a text-first DOCX and split it into chapters.")
    parser.add_argument("--source", default="data/source/book.docx")
    parser.add_argument("--clean-dir", default="build/clean")
    parser.add_argument("--text-dir", default="build/text")
    args = parser.parse_args()

    manifest = prepare(Path(args.source), Path(args.clean_dir), Path(args.text_dir))
    print(f"Prepared {len(manifest['chapters'])} chapter unit(s).")
    print(f"Clean DOCX: {manifest['clean_docx']}")
    print("Next: inspect the clean DOCX, then run Word -> DTBook in DAISY Pipeline 2.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
